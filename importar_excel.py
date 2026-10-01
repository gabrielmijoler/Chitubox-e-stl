"""
Importa registos do CSV (rodada do robô) para a planilha Excel de destino.

- Lê o CSV mais recente da pasta 'resultados/' (ou o definido em CSV_ORIGEM)
- Mapeia apenas 3 campos: Nome do Arquivo, Peso (g) e Tempo (h)
- Insere na primeira linha vazia da folha do Excel, sem sobrescrever dados
- Ignora registos já existentes (duplicados)
- Cria o Excel com cabeçalho se ainda não existir

Uso:
    python importar_excel.py                       # usa as constantes abaixo
    python importar_excel.py caminho/do.csv        # CSV à escolha
    python importar_excel.py caminho/do.csv caminho/do.xlsx
"""

from __future__ import annotations

import math
import re
import shutil
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

# --- CONFIGURAÇÃO (editável) ---
PASTA_BASE = Path(__file__).resolve().parent
PASTA_RESULTADOS = PASTA_BASE / "resultados"

# None = usa automaticamente o CSV mais recente da pasta PASTA_RESULTADOS
CSV_ORIGEM: str | None = None

EXCEL_DESTINO = r"D:\Arquivos Impressão 3D\Shopee\New Calculadora 3D - Copia.xlsx"
FOLHA_DESTINO = "Pecas"

# Posições das colunas na folha destino (A..E)
COL_NOME = 1    # A: Nome do arquivo
COL_PESO = 4    # D: Peso (g)
COL_TEMPO = 5   # E: Tempo (h)
CABECALHOS_NOVOS = ["Nome do arquivo", "Altura (mm)", "Largura (mm)", "Peso (g)", "Tempo (h)"]

# Sinónimos aceites por campo (cabeçalhos comparados em minúsculas, sem espaços extra)
SINONIMOS: dict[str, tuple[str, ...]] = {
    "nome": ("ficheiro", "file", "nome do arquivo", "nome do ficheiro", "arquivo"),
    "peso": ("weight_g", "weight", "peso (g)", "peso g", "peso"),
    "tempo": ("time", "tempo", "tempo (h)", "tempo h"),
}


# --- CSV ---

def resolver_csv_origem() -> Path:
    """Devolve o caminho do CSV a usar (constante ou o mais recente de resultados/)."""
    if CSV_ORIGEM:
        caminho = Path(CSV_ORIGEM)
        if not caminho.is_file():
            raise FileNotFoundError(f"CSV não encontrado: {caminho}")
        return caminho

    if not PASTA_RESULTADOS.is_dir():
        raise FileNotFoundError(f"Pasta de CSVs não encontrada: {PASTA_RESULTADOS}")

    candidatos = sorted(
        PASTA_RESULTADOS.glob("resultados_*.csv"),
        key=lambda p: p.stat().st_mtime,
    )
    if not candidatos:
        raise FileNotFoundError(f"Nenhum CSV 'resultados_*.csv' encontrado em {PASTA_RESULTADOS}")
    return candidatos[-1]


def ler_csv(caminho: Path) -> pd.DataFrame:
    """Lê o CSV e normaliza os cabeçalhos (minúsculas, sem espaços extra)."""
    try:
        df = pd.read_csv(caminho, dtype=str, encoding="utf-8-sig")
    except pd.errors.EmptyDataError as erro:
        raise ValueError(f"CSV vazio (sem dados): {caminho}") from erro

    df = df.dropna(how="all")
    df.columns = [re.sub(r"\s+", " ", str(c)).strip().lower() for c in df.columns]
    return df


def mapear_colunas(colunas: list[str]) -> dict[str, str]:
    """Associa cada campo obrigatório à coluna real do CSV (resiliente a maiúsculas/espaços)."""
    mapeado: dict[str, str] = {}
    for campo, opcoes in SINONIMOS.items():
        for coluna in colunas:
            if coluna in opcoes:
                mapeado[campo] = coluna
                break

    faltantes = [campo for campo in SINONIMOS if campo not in mapeado]
    if faltantes:
        raise ValueError(
            "Coluna(s) obrigatória(s) em falta no CSV: "
            + ", ".join(faltantes)
            + f". Colunas encontradas: {', '.join(colunas) or '(nenhuma)'}"
        )
    return mapeado


# --- CONVERSÕES ---

def limpar(valor) -> str:
    """Remove nulos e espaços em branco de um valor vindo do CSV."""
    if valor is None or (isinstance(valor, float) and math.isnan(valor)):
        return ""
    return str(valor).strip()


def converter_peso(valor) -> float | None:
    """'10.039 g' -> 10.039 ; aceita também vírgula decimal ('10,039')."""
    texto = limpar(valor)
    if not texto:
        return None
    if texto.startswith("="):  # fórmula do Excel (ex.: kits agregados) — não é um valor
        return None
    numeros = re.sub(r"[^\d.,-]", "", texto)
    if not numeros:
        return None
    if "," in numeros and "." in numeros:      # formato pt-BR: 1.234,56
        numeros = numeros.replace(".", "").replace(",", ".")
    elif "," in numeros:                        # vírgula decimal: 10,039
        numeros = numeros.replace(",", ".")
    try:
        return round(float(numeros), 3)
    except ValueError:
        return None


def converter_tempo(valor) -> float | None:
    """'3h2m34s' -> 3.04 (horas decimais com 2 casas); aceita também horas puras ('3.3')."""
    texto = limpar(valor).lower()
    if not texto:
        return None

    correspondencia = re.fullmatch(
        r"\s*(?:(\d+)\s*h)?\s*(?:(\d+)\s*m)?\s*(?:(\d+)\s*s)?\s*", texto
    )
    if correspondencia and any(correspondencia.groups()):
        horas, minutos, segundos = (
            int(g) if g else 0 for g in correspondencia.groups()
        )
        return round((horas * 3600 + minutos * 60 + segundos) / 3600, 2)

    try:
        return round(float(texto.replace(",", ".")), 2)
    except ValueError:
        return None


def extrair_registos(df: pd.DataFrame, colunas: dict[str, str]) -> tuple[list[tuple[str, float, float]], int]:
    """Converte as linhas do CSV em registos (nome, peso, tempo). Devolve (registos, ignoradas)."""
    registos: list[tuple[str, float, float]] = []
    ignoradas = 0

    for _, linha in df.iterrows():
        nome = limpar(linha[colunas["nome"]])
        peso = converter_peso(linha[colunas["peso"]])
        tempo = converter_tempo(linha[colunas["tempo"]])

        if not nome:
            ignoradas += 1
            print("  ⚠️ Linha ignorada: nome vazio.")
            continue
        if peso is None or tempo is None:
            ignoradas += 1
            print(f"  ⚠️ '{nome}': peso ou tempo inválido — linha ignorada.")
            continue

        registos.append((nome, peso, tempo))

    return registos, ignoradas


# --- EXCEL ---

def abrir_ou_criar_excel(caminho: Path) -> tuple[Workbook, bool]:
    """Abre o Excel de destino ou cria um novo com cabeçalho. Devolve (workbook, era_novo)."""
    if not caminho.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = FOLHA_DESTINO
        for indice, cabecalho in enumerate(CABECALHOS_NOVOS, start=1):
            ws.cell(row=1, column=indice, value=cabecalho)
        print(f"ℹ️ Excel não encontrado — criado novo com cabeçalho: {caminho}")
        return wb, True

    wb = load_workbook(caminho)
    if FOLHA_DESTINO not in wb.sheetnames:
        raise ValueError(
            f"Folha '{FOLHA_DESTINO}' não existe no Excel. Folhas disponíveis: {', '.join(wb.sheetnames)}"
        )
    return wb, False


def ler_linhas_existentes(ws) -> set[tuple[str, float | None, float | None]]:
    """Conjunto de chaves (nome, peso, tempo) já preenchidas, para ignorar duplicados."""
    existentes: set[tuple[str, float | None, float | None]] = set()
    for linha in range(2, ws.max_row + 1):
        nome = ws.cell(row=linha, column=COL_NOME).value
        if nome is None or str(nome).strip() == "":
            continue
        existentes.add((
            str(nome).strip(),
            converter_peso(ws.cell(row=linha, column=COL_PESO).value),
            converter_tempo(ws.cell(row=linha, column=COL_TEMPO).value),
        ))
    return existentes


def primeira_linha_vazia(ws) -> int:
    """Primeira linha (a partir da 2) sem 'Nome do arquivo' — nunca sobrescreve dados."""
    for linha in range(2, ws.max_row + 2):
        if ws.cell(row=linha, column=COL_NOME).value in (None, ""):
            return linha
    return ws.max_row + 1


def proxima_linha_livre(ws, linha: int) -> int:
    """Avança até uma linha sem 'Nome do arquivo' — nunca sobrescreve linhas ocupadas."""
    while ws.cell(row=linha, column=COL_NOME).value not in (None, ""):
        linha += 1
    return linha


def expandir_tabela(ws, ultima_linha: int) -> None:
    """Alarga a tabela/AutoFilter do Excel se escrevemos além do fim atual (ex.: A1:N2206)."""
    for tabela in ws.tables.values():
        if not isinstance(tabela, Table):
            continue
        inicio, fim = tabela.ref.split(":")
        coluna_fim = re.sub(r"\d", "", fim)
        linha_fim = int(re.sub(r"\D", "", fim))
        if ultima_linha > linha_fim:
            tabela.ref = f"{inicio}:{coluna_fim}{ultima_linha}"
            if tabela.autoFilter is not None:
                tabela.autoFilter.ref = tabela.ref
            print(f"  📐 Tabela expandida para {tabela.ref}")


# --- PRINCIPAL ---

def main() -> int:
    caminho_csv = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    caminho_excel = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(EXCEL_DESTINO)

    # 1) Resolver e ler o CSV
    try:
        if caminho_csv is None:
            caminho_csv = resolver_csv_origem()
        elif not caminho_csv.is_file():
            raise FileNotFoundError(f"CSV não encontrado: {caminho_csv}")

        df = ler_csv(caminho_csv)
        colunas = mapear_colunas(list(df.columns))
    except (FileNotFoundError, ValueError) as erro:
        print(f"❌ {erro}")
        return 1

    print(f"📄 CSV: {caminho_csv}")
    print(f"🔎 Colunas mapeadas: nome='{colunas['nome']}', peso='{colunas['peso']}', tempo='{colunas['tempo']}'")

    # 2) Converter as linhas em registos limpos
    registos, ignoradas = extrair_registos(df, colunas)
    if not registos:
        print(f"⚠️ Nenhum registo válido no CSV ({ignoradas} linha(s) ignorada(s)).")
        return 0

    # 3) Abrir/criar o Excel
    try:
        wb, era_novo = abrir_ou_criar_excel(caminho_excel)
    except Exception as erro:  # noqa: BLE001 — log claro para o utilizador
        print(f"❌ Erro ao abrir o Excel: {erro}")
        return 1

    ws = wb[FOLHA_DESTINO]
    existentes = ler_linhas_existentes(ws)
    linha_actual = primeira_linha_vazia(ws)

    print(f"📘 Excel: {caminho_excel} | Folha: '{FOLHA_DESTINO}' | 1ª linha vazia: {linha_actual}")

    # 4) Escrever A/D/E, ignorando duplicados e linhas ocupadas
    inseridos = 0
    duplicados = 0
    primeira_escrita: int | None = None
    ultima_escrita: int | None = None
    for nome, peso, tempo in registos:
        chave = (nome, peso, tempo)
        if chave in existentes:
            duplicados += 1
            print(f"  ⏭️ {nome}: já existe no Excel — ignorado.")
            continue

        linha_actual = proxima_linha_livre(ws, linha_actual)
        ws.cell(row=linha_actual, column=COL_NOME, value=nome)
        ws.cell(row=linha_actual, column=COL_PESO, value=peso)
        ws.cell(row=linha_actual, column=COL_TEMPO, value=tempo)
        print(f"  ✅ {nome}: {peso} g | {tempo} h → linha {linha_actual}")

        existentes.add(chave)
        if primeira_escrita is None:
            primeira_escrita = linha_actual
        ultima_escrita = linha_actual
        linha_actual += 1
        inseridos += 1

    # 5) Guardar apenas se houver novidades
    if inseridos == 0:
        print(f"\n📊 Importados: 0 | Duplicados ignorados: {duplicados} | Inválidas: {ignoradas}")
        print("Nada a gravar — o Excel não foi alterado.")
        return 0

    expandir_tabela(ws, ultima_escrita)
    wb.calculation.fullCalcOnLoad = True  # força o Excel a recalcular as fórmulas ao abrir

    # Backup antes de gravar (protege os dados já existentes)
    if not era_novo and caminho_excel.exists():
        carimbo = time.strftime("%Y-%m-%d_%H-%M-%S")
        backup = caminho_excel.with_name(f"{caminho_excel.stem}_backup_{carimbo}{caminho_excel.suffix}")
        shutil.copy2(caminho_excel, backup)
        print(f"💾 Backup criado: {backup.name}")

    try:
        wb.save(caminho_excel)
    except PermissionError:
        print(f"❌ Não foi possível gravar — o ficheiro está aberto no Excel? Feche-o e tente de novo: {caminho_excel}")
        return 1

    print(f"\n📊 Importados: {inseridos} | Duplicados ignorados: {duplicados} | Inválidas: {ignoradas}")
    print(f"📗 Gravado em: {caminho_excel} (folha '{FOLHA_DESTINO}', linhas {primeira_escrita}..{ultima_escrita})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
