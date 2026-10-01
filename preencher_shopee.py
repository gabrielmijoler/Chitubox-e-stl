#!/usr/bin/env python3
"""
Preenche o template de criação em massa da Shopee com as peças da
New Calculadora 3D (linhas 339-385 = 47 peças da migração).

Modo agrupado (AGRUPAR_E): as peças são agrupadas por décadas de preço
(int(preco//10)*10), lotes de até 10 variações, e cada grupo vira UM
anúncio com variações (coluna E = chave de integração, mesma em todo
o grupo). Título/descrição do grupo são genéricos (corrigíveis à mão);
fotos (H e R..Z) ficam vazias para preenchimento manual.

Fases:
  1. Preços  - Excel COM (read-only, ficheiro intacto): recalcular e ler K/L
  2. Linhas  - montar em memória (grupos, variações, preços, canais)
  3. Gravar  - patch activePane -> openpyxl -> validar listas -> salvar
  4. Validar  - reabrir ficheiro final (DV counts, diff, grupos, valores)

Uso: .\\.venv\\Scripts\\python.exe preencher_shopee.py
"""

from __future__ import annotations

import io
import re
import subprocess
import sys
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# --- CONFIGURAÇÃO ----------------------------------------------------------
CALCULADORA = Path(r"D:\Arquivos Impressão 3D\Shopee\New Calculadora 3D - Copia.xlsx")
TEMPLATE = Path(r"D:\Downloads\Shopee_mass_upload_2026-09-30_basic_template.xlsx")
SAIDA = Path(r"D:\Downloads\Shopee_mass_upload_2026-09-30_agrupado.xlsx")

FOLHA_CALC = "Pecas"
FOLHA_TEMPLATE = "Modelo"
LINHA_INI, LINHA_FIM = 339, 385          # bloco pós-último kit (339 = 1ª linha da migração)
LINHA_DADOS_TEMPLATE = 7                  # primeira linha de dados do template

# --- Agrupamento em variações (coluna E) ------------------------------------
AGRUPAR_E = True                          # const: False = 1 anúncio por peça (modo antigo)
GRUPO_MAX_VAR = 10                        # máx. de variações por anúncio; década cheia -> Lote novo
PRE_ENCOMENDA = 3                         # AG: dias de postagem p/ encomenda (faixa cat 101385 = 3..15)
NOME_VAR = "Modelo"                       # F: nome da variação 1

CATEGORIA = "101385"                      # Hobbies e Coleções/Itens Colecionáveis/Figuras de Ação
ESTOQUE = 99999
PESO_KG = 0.5                             # peso fixo do pacote (kg)
COMPRIMENTO, LARGURA, ALTURA = 6, 16, 8   # caixa 6x16x8 cm, ordem escrita pelo utilizador
CANAL_AE = "Ligado"                       # Shopee Xpress CPF - DV "Ligado,Desativado"
ITEM_AGRUPAVEL = "No"                     # DV "Yes,No" com allowBlank=false
GTIN_SEM = "00"                           # P (GTIN): "00" = item sem GTIN (texto)
# AF (Retirada pelo Comprador) fica VAZIA - template: "Please do not edit this column"

PREFIXO_TITULO = "Miniaturas RPG - "
SUFIXO_TITULO = " - D&D Pathfinder Tormenta 20"
LIMITE_TITULO = 120
LIMITE_G = 30                             # coluna G: 1 a 30 caracteres
LIMITE_M = 100                            # coluna M: menos de 100 caracteres
RAZAO_PRECO_MAX = 4.00                    # Shopee: preço mais caro / mais barato do anúncio

DESCRICAO_MODELO = (
    "ATENÇÃO\n"
    "\n"
    "Caso não encontre uma miniatura que necessite, ou deseje personalizar com uma "
    "pintura customizada, entre em contato pelo chat, que estaremos felizes em te auxiliar.\n"
    "\n"
    "  \n"
    "\n"
    "Dimensões aproximadas das peças \n"
    "\n"
    "@NOME@ | @NOME@\n"
    "\n"
    "\n"
    "IMPORTANTE: Produto enviado com base separado do restante da peça e com primer "
    "preto, impressão é tratada com luz UV.\n"
    "\n"
    "\n"
    "Somos uma marca preocupada em entregar o melhor produto para todos compradores. "
    "Nosso foco está em miniaturas de RPG, criação de personagens customizados, cenários "
    "e terrenos para enriquecer a experiência de suas mesas, seja em campos de batalhas, "
    "praças das cidades, esconderijos de ladrões, hordas de zumbis, monstros aterrorizadores "
    "e mercadores em apuros."
)

# Colunas do template (Modelo)
COL = {
    "categoria": 1, "nome": 2, "descricao": 3, "sku_principal": 4,
    "e_agrup": 5, "f_var": 6, "g_opcao": 7, "preco": 11, "estoque": 12,
    "m_sku": 13, "gtin": 16, "peso": 27, "comprimento": 28, "largura": 29,
    "altura": 30, "canal_ae": 31, "ag": 33, "agrupavel": 48,
}

# --- FASE 1: preços via Excel COM ------------------------------------------
PS_TEMPLATE = """
$ErrorActionPreference = 'Stop'
$planilha = '@@PLANILHA@@'
$linhaIni = @@INI@@
$linhaFim = @@FIM@@
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
$excel.ScreenUpdating = $false
try {
    $wb = $excel.Workbooks.Open($planilha, 0, $true)
    $excel.CalculateFullRebuild()
    $ws = $wb.Worksheets.Item('Pecas')
    $inv = [System.Globalization.CultureInfo]::InvariantCulture
    for ($r = $linhaIni; $r -le $linhaFim; $r++) {
        $k = $ws.Cells.Item($r, 11).Value2
        $l = $ws.Cells.Item($r, 12).Value2
        $ks = if ($null -eq $k -or $k -is [string]) { '' } else { ([double]$k).ToString('R', $inv) }
        $ls = if ($null -eq $l -or $l -is [string]) { '' } else { ([double]$l).ToString('R', $inv) }
        Write-Output ('{0}|{1}|{2}' -f $r, $ks, $ls)
    }
    $wb.Close($false)
} finally {
    $excel.Quit()
    [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel)
}
"""


def obter_precos_com() -> dict[int, tuple[float | None, float | None]]:
    """Abre a New Calculadora read-only, recalcula e devolve {linha: (K, L)}."""
    ps = (
        PS_TEMPLATE
        .replace("@@PLANILHA@@", str(CALCULADORA))
        .replace("@@INI@@", str(LINHA_INI))
        .replace("@@FIM@@", str(LINHA_FIM))
    )
    print(f"  Excel COM: abrir read-only, CalculateFullRebuild, ler K/L {LINHA_INI}..{LINHA_FIM} ...", flush=True)
    proc = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=300,
    )
    if proc.returncode != 0:
        raise RuntimeError(
            "Excel COM falhou (verifica se a Excel está licenciada e a planilha fechada):\n"
            + (proc.stderr or proc.stdout)
        )

    dados: dict[int, tuple[float | None, float | None]] = {}
    for linha in proc.stdout.splitlines():
        m = re.match(r"^(\d+)\|([^|]*)\|(.*)$", linha.strip())
        if not m:
            continue
        r = int(m.group(1))
        k = float(m.group(2)) if m.group(2).strip() else None
        l = float(m.group(3)) if m.group(3).strip() else None
        dados[r] = (k, l)

    esperadas = LINHA_FIM - LINHA_INI + 1
    if len(dados) != esperadas:
        raise RuntimeError(f"COM devolveu {len(dados)} linhas (esperado {esperadas}). stdout:\n{proc.stdout}")
    return dados


# --- FASE 2: montar linhas ---------------------------------------------------
def limpar_nome_titulo(nome: str) -> str:
    n = re.sub(r"\s+", " ", str(nome).replace("_", " ")).strip()
    return n.strip("| ").strip()


def construir_titulo(nome: str) -> str:
    nome_t = limpar_nome_titulo(nome)
    max_nome = LIMITE_TITULO - len(PREFIXO_TITULO) - len(SUFIXO_TITULO)
    if len(nome_t) > max_nome:
        nome_t = nome_t[:max_nome].rstrip()
    titulo = PREFIXO_TITULO + nome_t + SUFIXO_TITULO
    if not 2 <= len(titulo) <= LIMITE_TITULO:
        raise ValueError(f"título inválido ({len(titulo)} car.): {titulo}")
    return titulo


def construir_descricao(nome: str) -> str:
    nome_t = limpar_nome_titulo(nome)
    desc = DESCRICAO_MODELO.replace("@NOME@", nome_t)
    if not 10 <= len(desc) <= 5000:
        raise ValueError(f"descrição inválida ({len(desc)} car.)")
    return desc


def preco_shopee(k: float | None, l: float | None, r: int) -> float:
    candidatos = [v for v in (k, l) if v is not None]
    if not candidatos:
        raise ValueError(f"linha {r}: K e L vazios - sem preço")
    p = Decimal(str(max(candidatos))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if not Decimal("1.00") <= p <= Decimal("100000.00"):
        raise ValueError(f"linha {r}: preço {p} fora de 1.00..100000.00")
    return float(p)


# --- FASE 2b: agrupamento por décadas de preço --------------------------------
def rotulo_decada(decada: int) -> str:
    return f"Peças {decada} a {decada + 9} Reais"


def construir_titulo_grupo(decada: int, lote: int, n_lotes: int) -> str:
    meio = rotulo_decada(decada)
    if n_lotes > 1:
        meio += f" Lote {lote}"
    max_nome = LIMITE_TITULO - len(PREFIXO_TITULO) - len(SUFIXO_TITULO)
    meio = meio[:max_nome].rstrip()
    titulo = PREFIXO_TITULO + meio + SUFIXO_TITULO
    if not 2 <= len(titulo) <= LIMITE_TITULO:
        raise ValueError(f"título de grupo inválido ({len(titulo)} car.): {titulo}")
    return titulo


def slug_sku(nome: str, prefixo: str = "MINI.") -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", ".", limpar_nome_titulo(nome)).strip(".").upper()
    sku = prefixo + slug
    if not 1 <= len(sku) <= LIMITE_M:
        raise ValueError(f"SKU de variação inválido ({len(sku)} car.): {sku}")
    return sku


def opcao_variacao(nome: str, usadas: set[str]) -> str:
    """G = nome da peça limpo, ≤30 chars, único dentro do grupo."""
    g = limpar_nome_titulo(nome)[:LIMITE_G].rstrip()
    base, n = g, 2
    while g in usadas:
        sufixo = f" {n}"
        g = base[:LIMITE_G - len(sufixo)].rstrip() + sufixo
        n += 1
    if not 1 <= len(g) <= LIMITE_G:
        raise ValueError(f"opção de variação inválida: {g!r}")
    return g


def agrupar_por_decada(precos: dict[int, float]) -> list[dict]:
    """Agrupa linhas por década de preço (int(p//10)*10), lotes de até GRUPO_MAX_VAR.

    Devolve lista de grupos: {decada, lote, n_lotes, linhas: [row...]}
    ordenados por preço crescente dentro de cada lote.
    """
    por_decada: dict[int, list[int]] = {}
    for r, p in precos.items():
        por_decada.setdefault(int(p // 10) * 10, []).append(r)

    grupos: list[dict] = []
    for decada in sorted(por_decada):
        linhas_d = sorted(por_decada[decada], key=lambda r: precos[r])
        lotes = [linhas_d[i:i + GRUPO_MAX_VAR] for i in range(0, len(linhas_d), GRUPO_MAX_VAR)]
        for i, lote in enumerate(lotes, start=1):
            grupos.append({
                "decada": decada, "lote": i, "n_lotes": len(lotes), "linhas": lote,
            })
    return grupos


# --- FASE 3: template + validação Gap 3 ---------------------------------------
def carregar_template() -> tuple:
    """Lê o ZIP do template, corrige activePane em memória e abre com openpyxl."""
    zin = zipfile.ZipFile(TEMPLATE)
    buf = io.BytesIO()
    corrigidos = 0
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                txt = dados.decode("utf-8")
                if 'activePane="bottom_left"' in txt:
                    txt = txt.replace('activePane="bottom_left"', 'activePane="bottomLeft"')
                    corrigidos += 1
                    dados = txt.encode("utf-8")
            zout.writestr(item, dados)
    zin.close()
    buf.seek(0)
    wb = load_workbook(buf)
    print(f"  Template aberto (activePane corrigido em {corrigidos} folhas)")
    return wb


def resolver_formula1(wb, formula1: str) -> list[str] | None:
    f = (formula1 or "").strip()
    if len(f) >= 2 and f.startswith('"') and f.endswith('"'):
        return [v.strip() for v in f[1:-1].split(",")]
    m = re.match(r"^'?([^'!]+)'?!(\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)$", f)
    if m:
        folha, ref1, ref2 = m.groups()
        folha = folha.replace("$", "")
        if folha not in wb.sheetnames:
            return None
        ws2 = wb[folha]
        intervalo = f"{ref1.replace('$', '')}:{ref2.replace('$', '')}"
        valores = []
        for linha in ws2[intervalo]:
            for cel in linha:
                if cel.value not in (None, ""):
                    valores.append(str(cel.value))
        return valores
    return None


def validar_contra_listas_da_planilha(ws, linha_final: int, linha_inicial: int = 7) -> list[str]:
    """Confere valores das colunas com DV type='list' (Gap 3). Devolve problemas."""
    wb = ws.parent
    problemas: list[str] = []
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        permitidos = resolver_formula1(wb, dv.formula1)
        if permitidos is None:
            problemas.append(f"DV list não interpretável: sqref={dv.sqref} formula1={dv.formula1}")
            continue
        set_perm = set(permitidos)
        for token in str(dv.sqref).split():
            min_c, min_r, max_c, max_r = range_boundaries(token)
            for r in range(max(min_r, linha_inicial), min(max_r, linha_final) + 1):
                for c in range(min_c, max_c + 1):
                    v = ws.cell(row=r, column=c).value
                    if v in (None, ""):
                        continue
                    if str(v).strip() not in set_perm:
                        problemas.append(
                            f"linha {r}, coluna {get_column_letter(c)}: valor {v!r} "
                            f"fora da lista (ex.: {sorted(set_perm)[:3]})"
                        )
    return problemas


def dv_set(wb) -> set:
    s = set()
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            s.add((ws.title, str(dv.sqref), str(dv.formula1), str(dv.type), bool(dv.allowBlank)))
    return s


def contar_dv_xml(caminho: Path) -> tuple[int, int, int]:
    z = zipfile.ZipFile(caminho)
    total = nlist = nwhole = 0
    for nome in z.namelist():
        if not nome.startswith("xl/worksheets/"):
            continue
        xml = z.read(nome).decode("utf-8", "replace")
        for tag in re.findall(r"<dataValidation\b[^>]*?>", xml):
            total += 1
            if 'type="list"' in tag:
                nlist += 1
            elif 'type="whole"' in tag:
                nwhole += 1
    z.close()
    return total, nlist, nwhole


# --- MAIN ---------------------------------------------------------------------
def main() -> int:
    print("=== FASE 1 - preços (Excel COM, read-only) ===")
    try:
        dados = obter_precos_com()
    except Exception as erro:  # noqa: BLE001
        print(f"❌ {erro}")
        return 1

    print("\n=== FASE 2 - montar linhas ===")
    # Nomes direto da calculadora (openpyxl; só a coluna A, não toca no ficheiro)
    wb_calc = load_workbook(CALCULADORA, read_only=True, data_only=True)
    ws_calc = wb_calc[FOLHA_CALC]
    nomes = {}
    for row in ws_calc.iter_rows(min_row=LINHA_INI, max_row=LINHA_FIM, min_col=1, max_col=1):
        nomes[row[0].row] = row[0].value
    wb_calc.close()

    pecas: list[dict] = []
    try:
        for r in range(LINHA_INI, LINHA_FIM + 1):
            nome, k, l = nomes.get(r), dados[r][0], dados[r][1]
            if not nome:
                raise ValueError(f"linha {r}: nome vazio na calculadora")
            pecas.append({"row": r, "nome": nome, "preco": preco_shopee(k, l, r)})
    except ValueError as erro:
        print(f"❌ {erro}")
        return 1

    precos = {p["row"]: p["preco"] for p in pecas}
    print(f"  {len(pecas)} peças | preço min {min(precos.values()):.2f} "
          f"| médio {sum(precos.values())/len(precos):.2f} | max {max(precos.values()):.2f}")

    # Agrupar em anúncios com variações (coluna E = chave de integração)
    if AGRUPAR_E:
        grupos = agrupar_por_decada(precos)
        print(f"  Agrupamento por décadas: {len(grupos)} anúncios "
              f"(máx {GRUPO_MAX_VAR} variações/grupo)")
    else:
        grupos = [{"decada": None, "lote": 1, "n_lotes": 1, "linhas": [p["row"] for p in pecas]}]

    por_row = {p["row"]: p for p in pecas}
    linhas: list[dict] = []          # uma entrada por LIGAÇA (linhas do template)
    grupos_out: list[dict] = []      # metadados p/ validação da Fase 4
    try:
        for e, g in enumerate(grupos, start=1):
            decada, lote, n_lotes = g["decada"], g["lote"], g["n_lotes"]
            if AGRUPAR_E:
                rotulo = rotulo_decada(decada) + (f" Lote {lote}" if n_lotes > 1 else "")
                titulo = construir_titulo_grupo(decada, lote, n_lotes)
                descricao = construir_descricao(rotulo)
                sku_principal = f"MINI.LOTE{decada}.{lote}"
                e_valor = str(e)
            else:
                rotulo = None
                titulo = descricao = sku_principal = e_valor = None

            usadas_g: set[str] = set()
            usadas_m: set[str] = set()
            item_grupo = {"e": e_valor, "titulo": titulo, "descricao": descricao,
                          "sku_principal": sku_principal, "linhas": []}
            for r in g["linhas"]:
                p = por_row[r]
                if AGRUPAR_E:
                    g_opcao = opcao_variacao(p["nome"], usadas_g)
                    usadas_g.add(g_opcao)
                    m_sku = slug_sku(p["nome"])
                    if m_sku in usadas_m:            # unicidade de M dentro do grupo
                        raise ValueError(f"linha {r}: SKU duplicado no grupo: {m_sku}")
                    usadas_m.add(m_sku)
                else:
                    titulo = construir_titulo(p["nome"])
                    descricao = construir_descricao(p["nome"])
                    g_opcao = m_sku = None
                item = {
                    "row": r, "nome": p["nome"], "preco": p["preco"],
                    "titulo": titulo, "descricao": descricao,
                    "sku_principal": sku_principal,
                    "e": e_valor, "g": g_opcao, "m": m_sku,
                }
                linhas.append(item)
                item_grupo["linhas"].append(item)
            grupos_out.append(item_grupo)
    except ValueError as erro:
        print(f"❌ {erro}")
        return 1

    print("\n=== FASE 3 - preencher template + validar listas ===")
    try:
        wb = carregar_template()
    except Exception as erro:  # noqa: BLE001
        print(f"❌Erro ao abrir template: {erro}")
        return 1
    ws = wb[FOLHA_TEMPLATE]

    for i, item in enumerate(linhas):
        r = LINHA_DADOS_TEMPLATE + i
        ws.cell(row=r, column=COL["categoria"], value=CATEGORIA)
        ws.cell(row=r, column=COL["nome"], value=item["titulo"])
        ws.cell(row=r, column=COL["descricao"], value=item["descricao"])
        ws.cell(row=r, column=COL["preco"], value=item["preco"])
        ws.cell(row=r, column=COL["estoque"], value=ESTOQUE)
        ws.cell(row=r, column=COL["peso"], value=PESO_KG)
        ws.cell(row=r, column=COL["comprimento"], value=COMPRIMENTO)
        ws.cell(row=r, column=COL["largura"], value=LARGURA)
        ws.cell(row=r, column=COL["altura"], value=ALTURA)
        ws.cell(row=r, column=COL["canal_ae"], value=CANAL_AE)
        ws.cell(row=r, column=COL["agrupavel"], value=ITEM_AGRUPAVEL)
        c_gtin = ws.cell(row=r, column=COL["gtin"], value=GTIN_SEM)
        c_gtin.number_format = "@"               # P: texto ("00" não vira número)
        # AF (32) fica vazia - "Please do not edit this column"
        if AGRUPAR_E:
            ws.cell(row=r, column=COL["sku_principal"], value=item["sku_principal"])
            ws.cell(row=r, column=COL["e_agrup"], value=item["e"])
            ws.cell(row=r, column=COL["f_var"], value=NOME_VAR)
            ws.cell(row=r, column=COL["g_opcao"], value=item["g"])
            ws.cell(row=r, column=COL["m_sku"], value=item["m"])
            ws.cell(row=r, column=COL["ag"], value=PRE_ENCOMENDA)
            # H (8) e R..Z: fotos - ficam vazias (preenchimento manual)

    linha_final = LINHA_DADOS_TEMPLATE + len(linhas) - 1
    problemas = validar_contra_listas_da_planilha(ws, linha_final=linha_final)
    if problemas:
        print(f"❌ VALIDAÇÃO FALHOU ({len(problemas)}) - nada gravado:")
        for p in problemas[:40]:
            print("   ", p)
        return 1
    print(f"  Validação das listas: OK (linhas {LINHA_DADOS_TEMPLATE}..{linha_final})")

    wb.save(SAIDA)
    wb.close()
    print(f"  Gravado: {SAIDA}")

    print("\n=== FASE 4 - validação pós-gravação ===")
    contagem = contar_dv_xml(SAIDA)
    esperado = (5005, 5004, 1)
    ok = True
    if contagem != esperado:
        print(f"  ❌ DVs no ficheiro gravado: {contagem} (esperado {esperado})")
        ok = False
    else:
        print(f"  DVs no XML gravado: {contagem} (total/list/whole) ✔")

    wb_t = carregar_template()
    wb_s = load_workbook(SAIDA)
    s_t, s_s = dv_set(wb_t), dv_set(wb_s)
    diff = s_t ^ s_s
    if diff:
        ok = False
        print(f"  ❌ diff de DVs ({len(diff)}):")
        for d in sorted(diff)[:20]:
            print("   ", d)
    else:
        print(f"  DVs (sqref+formula1+type+allowBlank): diff = 0 ✔")

    ws_o, ws_f = wb_t[FOLHA_TEMPLATE], wb_s[FOLHA_TEMPLATE]
    alteradas = []
    for r in range(1, 7):
        for c in range(1, 53):
            a = ws_o.cell(row=r, column=c).value
            b = ws_f.cell(row=r, column=c).value
            # openpyxl normaliza '' -> None na gravação; ambos contam como vazio
            if (a if a not in (None, "") else None) != (b if b not in (None, "") else None):
                alteradas.append(f"{get_column_letter(c)}{r}")
    if alteradas:
        ok = False
        print(f"  ❌ linhas 1-6 alteradas: {alteradas}")
    else:
        print("  Linhas 1-6 intactas ✔")

    problemas_f = validar_contra_listas_da_planilha(ws_f, linha_final=linha_final)
    if problemas_f:
        ok = False
        print(f"  ❌ revalidação do ficheiro final: {len(problemas_f)} problemas")
        for p in problemas_f[:20]:
            print("   ", p)
    else:
        print("  Revalidação listas no ficheiro final: OK ✔")

    erros_valores = []
    for i, item in enumerate(linhas):
        r = LINHA_DADOS_TEMPLATE + i
        esperados = {
            COL["categoria"]: CATEGORIA, COL["nome"]: item["titulo"],
            COL["descricao"]: item["descricao"], COL["preco"]: item["preco"],
            COL["estoque"]: ESTOQUE, COL["peso"]: PESO_KG,
            COL["comprimento"]: COMPRIMENTO, COL["largura"]: LARGURA,
            COL["altura"]: ALTURA, COL["canal_ae"]: CANAL_AE,
            COL["agrupavel"]: ITEM_AGRUPAVEL, 32: None,
            COL["gtin"]: GTIN_SEM,
        }
        if AGRUPAR_E:
            esperados.update({
                COL["sku_principal"]: item["sku_principal"], COL["e_agrup"]: item["e"],
                COL["f_var"]: NOME_VAR, COL["g_opcao"]: item["g"],
                COL["m_sku"]: item["m"], COL["ag"]: PRE_ENCOMENDA,
                34: None,                   # AH: sem NCM
            })
        for c, esp in esperados.items():
            v = ws_f.cell(row=r, column=c).value
            if v != esp:
                erros_valores.append(f"{get_column_letter(c)}{r}: {v!r} != {esp!r}")
    if ws_f.cell(row=linha_final + 1, column=COL["nome"]).value not in (None, ""):
        erros_valores.append(f"B{linha_final + 1} devia estar vazia")
    if erros_valores:
        ok = False
        print(f"  ❌ valores divergentes ({len(erros_valores)}):")
        for e in erros_valores[:20]:
            print("   ", e)
    else:
        print(f"  Valores das {len(linhas)} linhas conferem ✔ "
              "(incl. AE=Ligado, AF=vazia, AV=No, P=00)")

    # --- validação dos grupos (coluna E) -------------------------------------
    if AGRUPAR_E:
        erros_grupo = []
        por_e: dict[str, list[dict]] = {}
        for item in linhas:
            por_e.setdefault(item["e"], []).append(item)
        if len(por_e) != len(grupos_out):
            erros_grupo.append(f"{len(por_e)} grupos distintos em E (esperado {len(grupos_out)})")
        for e, itens in por_e.items():
            if len(itens) > GRUPO_MAX_VAR:
                erros_grupo.append(f"E={e}: {len(itens)} variações (máx {GRUPO_MAX_VAR})")
            for campo in ("titulo", "descricao", "sku_principal"):
                if len({it[campo] for it in itens}) != 1:
                    erros_grupo.append(f"E={e}: {campo} divergente dentro do grupo")
            gs = [it["g"] for it in itens]
            ms = [it["m"] for it in itens]
            if len(set(gs)) != len(gs):
                erros_grupo.append(f"E={e}: opção G duplicada")
            if len(set(ms)) != len(ms):
                erros_grupo.append(f"E={e}: SKU M duplicado")
            ps = [it["preco"] for it in itens]
            razao = max(ps) / min(ps)
            if razao > RAZAO_PRECO_MAX:
                erros_grupo.append(f"E={e}: razão preço {razao:.2f} > {RAZAO_PRECO_MAX}")
            # linhas do mesmo grupo contíguas no template?
            idxs = [linhas.index(it) for it in itens]
            if idxs != list(range(idxs[0], idxs[0] + len(idxs))):
                erros_grupo.append(f"E={e}: linhas não contíguas no template")
        if erros_grupo:
            ok = False
            print(f"  ❌ validação de grupos ({len(erros_grupo)}):")
            for g in erros_grupo[:20]:
                print("   ", g)
        else:
            print(f"  Grupos: {len(por_e)} anúncios, ≤{GRUPO_MAX_VAR} var., B/C/D iguais, "
                  f"G/M únicos, razão preço ≤{RAZAO_PRECO_MAX} ✔")

        # fotos: H (por variação) e R..Z (galeria) ficam vazias - aviso, não falha
        cols_foto = [8] + list(range(18, 27))
        fotos_pendentes = [
            f"{get_column_letter(c)}{LINHA_DADOS_TEMPLATE + i}"
            for i in range(len(linhas)) for c in cols_foto
            if ws_f.cell(row=LINHA_DADOS_TEMPLATE + i, column=c).value in (None, "")
        ]
        if fotos_pendentes:
            print(f"  ⚠ FOTOS PENDENTES: {len(fotos_pendentes)} células vazias "
                  f"(H por variação + R..Z capa/galeria) - preencher manualmente antes do upload")

    wb_t.close()
    wb_s.close()

    print("\n=== AMOSTRA (linha template | preço | E | opção G | título) ===")
    for i in list(range(3)) + list(range(len(linhas) - 2, len(linhas))):
        item = linhas[i]
        e = item["e"] or "-"
        g = item["g"] or item["nome"]
        print(f"  {LINHA_DADOS_TEMPLATE + i} | {item['preco']:.2f} | {e} | {g} | {item['titulo']}")

    print()
    if ok:
        n_grupos = len(grupos_out) if AGRUPAR_E else len(linhas)
        print(f"✅ TUDO OK - ficheiro pronto para upload: {SAIDA}")
        print(f"   {len(linhas)} linhas = {n_grupos} anúncio(s) | linhas "
              f"{LINHA_DADOS_TEMPLATE}..{linha_final} do template")
        if AGRUPAR_E:
            print("   Pendente manual: fotos (H + R..Z), corrigir títulos genéricos,")
            print("   apagar os 47 anúncios antigos (GTIN já = 00; confirmar se o")
            print("   checkbox 'Item sem GTIN' ainda é necessário no upload)")
        return 0
    print("❌ FASE 4 FALHOU - rever erros acima")
    return 1


if __name__ == "__main__":
    sys.exit(main())
